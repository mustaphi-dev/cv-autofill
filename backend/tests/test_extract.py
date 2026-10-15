import io
import sys
import time
import zipfile
import zlib
from pathlib import Path

import docx
import pytest
from docx.oxml import parse_xml
from reportlab.lib.pagesizes import A4
from reportlab.lib.pdfencrypt import StandardEncryption
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import SimpleDocTemplate, Table

from app import extract
from app.extract import Unreadable
from app.extract import extract as read

# Synthetic CVs for a made-up person, built here so the repo holds no binaries
# and no real CV. Real CVs are checked locally from backend/local-cvs/.
WIDTH, HEIGHT = A4
LOREM = "Shipped the thing that made the other thing faster for everyone. " * 2


def pdf(draw, **options) -> bytes:  # type: ignore[no-untyped-def]
    out = io.BytesIO()
    canvas = Canvas(out, pagesize=A4, **options)
    draw(canvas)
    canvas.save()
    return out.getvalue()


def two_columns(c: Canvas) -> None:
    c.setFont("Helvetica-Bold", 18)
    c.drawString(40, HEIGHT - 50, "Jane Example, Senior Engineer and Team Lead")
    c.setFont("Helvetica", 9)
    c.linkURL(
        "https://linkedin.com/in/jane-example", (40, HEIGHT - 80, 120, HEIGHT - 66)
    )
    c.drawString(40, HEIGHT - 76, "LinkedIn")
    # A sidebar and a main column, each with its own line spacing.
    sidebar = ["SKILLS", "Python", "Rust", "SQL", "LANGUAGES", "English", "French"]
    main = ["EXPERIENCE", "Engineer at Acme", f"• {LOREM[:60]}", f"• {LOREM[:58]}"]
    main += [
        "Developer at Initech",
        f"• {LOREM[:55]}",
        "EDUCATION",
        "BSc, Uni of Leeds",
    ]
    for i, text in enumerate(sidebar):
        c.drawString(40, HEIGHT - 120 - i * 19, text)
    for i, text in enumerate(main):
        c.drawString(200, HEIGHT - 120 - i * 14, text)


def single_column(c: Canvas) -> None:
    # Right-aligned dates on the title row: a gap, but not two columns.
    y = HEIGHT - 50
    for title, dates in [
        ("Engineer at Acme", "Jan 2024 – Present"),
        ("Developer at Initech", "2020 – 2023"),
    ]:
        c.drawString(40, y, title)
        c.drawRightString(WIDTH - 40, y, dates)
        for _ in range(3):
            y -= 14
            c.drawString(40, y, f"• {LOREM[:100]}")
        y -= 20


def test_two_columns_read_header_then_left_then_right() -> None:
    text = read(pdf(two_columns))
    lines = text.splitlines()
    assert lines[0] == "Jane Example, Senior Engineer and Team Lead"
    assert lines.index("French") < lines.index("EXPERIENCE")
    assert lines[lines.index("EXPERIENCE") + 1] == "Engineer at Acme"
    # The address behind "LinkedIn" is recovered.
    assert text.endswith("Links:\nhttps://linkedin.com/in/jane-example")


def test_single_column_with_right_aligned_dates_is_not_split() -> None:
    lines = read(pdf(single_column)).splitlines()
    assert lines[0] == "Engineer at Acme | Jan 2024 – Present"
    assert lines[1].startswith("• Shipped")
    assert lines[4] == "Developer at Initech | 2020 – 2023"


def test_table_rows_stay_together() -> None:
    out = io.BytesIO()
    rows = [["Email", "jane@example.com"], ["Phone", "+44 20 7946 0000"]] * 5
    rows += [["Summary", LOREM[:70]]]
    SimpleDocTemplate(out).build([Table(rows, colWidths=[120, 300])])
    assert "Email | jane@example.com" in read(out.getvalue()).splitlines()


def test_unmapped_bullet_glyphs_become_bullets() -> None:
    row = [{"text": "(cid:127)", "x0": 0, "x1": 4}, {"text": "Led", "x0": 8, "x1": 20}]
    assert extract._line(row) == "• Led"
    assert extract._line([{"text": "\uf0b7Led", "x0": 0, "x1": 20}]) == "• Led"


def test_a_scanned_pdf_is_refused() -> None:
    with pytest.raises(Unreadable, match="scanned"):
        read(pdf(lambda c: c.rect(40, 40, 400, 600, fill=1)))


def test_a_password_protected_pdf_is_refused() -> None:
    locked = pdf(two_columns, encrypt=StandardEncryption("secret"))
    with pytest.raises(Unreadable, match="password"):
        read(locked)


def test_over_the_page_limit_is_refused() -> None:
    def pages(c: Canvas) -> None:
        for _ in range(extract.MAX_PAGES + 1):
            c.drawString(40, 400, LOREM)
            c.showPage()

    with pytest.raises(Unreadable, match="pages"):
        read(pdf(pages))


@pytest.mark.parametrize(
    ("data", "message"),
    [
        (b"plain text", extract.NOT_SUPPORTED),
        (b"\xd0\xcf\x11\xe0 old .doc", extract.NOT_SUPPORTED),
        (b"%PDF-1.7 broken", extract.UNREADABLE),
        (b"%PDF-" + b"0" * extract.MAX_BYTES, extract.TOO_BIG),
    ],
    ids=["text", "doc", "broken pdf", "too big"],
)
def test_refuses_other_files(data: bytes, message: str) -> None:
    with pytest.raises(Unreadable) as refused:
        read(data)
    assert str(refused.value) == message


def test_a_zip_that_is_not_a_docx_is_refused() -> None:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w") as archive:
        archive.writestr("notes.txt", "hello")
    with pytest.raises(Unreadable, match="PDF or DOCX"):
        read(out.getvalue())


def a_docx() -> bytes:
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Jane Example"
    document.add_paragraph("EXPERIENCE")
    document.add_paragraph("Shipped the thing", style="List Bullet")
    data = document.add_table(rows=2, cols=2)
    for row, (label, value) in zip(
        data.rows, [("Email", "jane@example.com"), ("Phone", "0123")], strict=True
    ):
        row.cells[0].text, row.cells[1].text = label, value
    layout = document.add_table(rows=1, cols=2)
    layout.cell(0, 0).text = "SKILLS"
    layout.cell(0, 0).add_paragraph("Python")
    layout.cell(0, 1).text = "PROJECTS"
    layout.cell(0, 1).add_paragraph("Robot", style="List Bullet")
    # A text box as Word saves it: a modern copy plus a fallback copy.
    text = "<w:p><w:r><w:t>In a box</w:t></w:r></w:p>"
    box = f"<w:p><w:r><w:txbxContent>{text}</w:txbxContent></w:r></w:p>"
    w = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
    mc = 'xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006"'
    document.element.body.insert(
        -1,  # Before the section settings, which stay last.
        parse_xml(
            f'<w:p {w} {mc}><w:r><mc:AlternateContent><mc:Choice Requires="wps">{box}'
            f"</mc:Choice><mc:Fallback>{box}</mc:Fallback></mc:AlternateContent></w:r></w:p>"
        ),
    )
    document.add_paragraph().add_run("LinkedIn")
    document.part.relate_to(
        "https://linkedin.com/in/jane-example",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def test_docx_reads_header_bullets_tables_and_text_boxes_in_order() -> None:
    assert read(a_docx()).splitlines() == [
        "Jane Example",
        "EXPERIENCE",
        "• Shipped the thing",
        "Email | jane@example.com",
        "Phone | 0123",
        "SKILLS",
        "Python",
        "PROJECTS",
        "• Robot",
        "In a box",
        "LinkedIn",
        "",
        "Links:",
        "https://linkedin.com/in/jane-example",
    ]


def test_docx_keeps_repeated_lines_but_not_repeated_headers() -> None:
    document = docx.Document()
    document.sections[0].header.paragraphs[0].text = "Jane Example"
    for _ in range(2):
        document.add_paragraph("Software Engineer")
        document.add_section().header.is_linked_to_previous = False
    for section in document.sections[1:]:
        section.header.paragraphs[0].text = "Jane Example"
    out = io.BytesIO()
    document.save(out)
    assert read(out.getvalue()).splitlines() == [
        "Jane Example",
        "Software Engineer",
        "Software Engineer",
    ]


def test_a_docx_zip_bomb_is_refused(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract, "MAX_UNZIPPED", 1000)
    with pytest.raises(Unreadable):
        extract._read(a_docx())


def test_a_stuck_worker_is_killed(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract, "TIMEOUT", 0.001)
    start = time.monotonic()
    with pytest.raises(Unreadable, match="couldn't read"):
        read(pdf(two_columns))
    assert time.monotonic() - start < 5


@pytest.mark.skipif(sys.platform != "linux", reason="memory limits are Linux-only")
def test_a_memory_bomb_kills_only_its_worker() -> None:
    # 2 GB of zeros, compressed to a few MB, as one page's content.
    packer = zlib.compressobj(9)
    megabyte = b"0" * 1024**2
    bomb = b"".join(packer.compress(megabyte) for _ in range(2048)) + packer.flush()
    data = (
        b"%PDF-1.7\n1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj\n"
        b"2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj\n"
        b"3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 612 792]"
        b"/Contents 4 0 R>>endobj\n"
        b"4 0 obj<</Length "
        + str(len(bomb)).encode()
        + b"/Filter/FlateDecode>>stream\n"
        + bomb
        + b"\nendstream endobj\ntrailer<</Root 1 0 R>>\n%EOF"
    )
    with pytest.raises(Unreadable):
        read(data)
    assert "Engineer at Acme" in read(pdf(two_columns))  # The next file works.


def test_nothing_is_written_to_disk(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr("tempfile.tempdir", str(tmp_path))
    for data in (pdf(two_columns), pdf(single_column), a_docx()):
        extract._read(data)  # In this process, so the patches apply.
    assert list(tmp_path.iterdir()) == []
